"""
AI Document Service - Save AI-generated analyses as project documents.
Automatically saves PDFs and Word documents when AI analyses are run.
"""
import os
import uuid
import logging
from datetime import datetime
from io import BytesIO

logger = logging.getLogger(__name__)


def save_ai_analysis_pdf(project, analysis_type, content, user_id=None):
    """
    Save AI analysis result as PDF in project documents.

    Args:
        project: Project model instance
        analysis_type: Type of analysis (risks, report, matching, document_analysis)
        content: Dict or string with analysis content
        user_id: ID of user who ran the analysis

    Returns:
        ProjectDocument instance or None if failed
    """
    from flask import current_app
    from models import db, ProjectDocument

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import cm

        # Generate PDF content
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        story = []

        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1a237e'),
            spaceAfter=20
        )
        section_style = ParagraphStyle(
            'SectionTitle',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#1a237e'),
            spaceBefore=15,
            spaceAfter=10
        )

        # Title based on analysis type
        type_titles = {
            'risks': 'Análise de Riscos',
            'report': 'Relatório do Projeto',
            'matching': 'Matching com Editais',
            'document_analysis': 'Análise de Documento'
        }

        title = type_titles.get(analysis_type, 'Análise IA')
        story.append(Paragraph(f"{title} - {project.code}", title_style))
        story.append(Paragraph(f"Projeto: {project.title}", styles['Normal']))
        story.append(Paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        story.append(Spacer(1, 20))

        # Process content based on type
        if analysis_type == 'risks':
            _add_risks_content(story, content, styles, section_style)
        elif analysis_type == 'report':
            _add_report_content(story, content, styles, section_style)
        elif analysis_type == 'matching':
            _add_matching_content(story, content, styles, section_style)
        elif analysis_type == 'document_analysis':
            _add_document_analysis_content(story, content, styles, section_style)

        # Footer
        story.append(Spacer(1, 30))
        story.append(Paragraph(
            "<i>Documento gerado automaticamente pelo Orion PMO IA</i>",
            ParagraphStyle('Footer', parent=styles['Normal'], fontSize=9, textColor=colors.gray)
        ))

        doc.build(story)
        pdf_content = buffer.getvalue()

        # Save to file
        filename = f"{analysis_type}_{project.code}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        stored_filename = f"{uuid.uuid4().hex}.pdf"

        tenant_folder = os.path.join(
            current_app.config.get('UPLOAD_FOLDER', 'uploads'),
            str(project.tenant_id or 'global')
        )
        os.makedirs(tenant_folder, exist_ok=True)

        file_path = os.path.join(tenant_folder, stored_filename)
        with open(file_path, 'wb') as f:
            f.write(pdf_content)

        # Create document record
        doc_record = ProjectDocument(
            tenant_id=project.tenant_id,
            project_id=project.id,
            filename=filename,
            stored_filename=stored_filename,
            file_type='pdf',
            file_size=len(pdf_content),
            description=f"{title} gerado por IA em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            uploaded_by_id=user_id
        )

        db.session.add(doc_record)
        db.session.commit()

        logger.info(f"Saved AI analysis PDF: {filename} for project {project.code}")
        return doc_record

    except Exception as e:
        logger.error(f"Error saving AI analysis PDF: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return None


def save_proposal_word(project, call, proposal_content, user_id=None):
    """
    Save generated proposal as Word document in project documents.

    Args:
        project: Project model instance
        call: PublicCall model instance
        proposal_content: Markdown/text content of the proposal
        user_id: ID of user who generated the proposal

    Returns:
        ProjectDocument instance or None if failed
    """
    from flask import current_app
    from models import db, ProjectDocument

    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title = doc.add_heading(f'Proposta - {project.title}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        doc.add_paragraph(f"Projeto: {project.code} - {project.title}")
        doc.add_paragraph(f"Edital: {call.source} - {call.title}")
        doc.add_paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        doc.add_paragraph()

        # Parse markdown content and add to document
        lines = proposal_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Handle headers
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
                doc.add_paragraph(line[3:], style='List Number')
            else:
                # Regular paragraph - remove markdown bold/italic
                clean_line = line.replace('**', '').replace('*', '').replace('__', '').replace('_', '')
                doc.add_paragraph(clean_line)

        # Footer
        doc.add_paragraph()
        footer = doc.add_paragraph("Documento gerado automaticamente pelo Orion PMO IA")
        footer.italic = True

        # Save to bytes
        buffer = BytesIO()
        doc.save(buffer)
        docx_content = buffer.getvalue()

        # Save to file
        filename = f"proposta_{project.code}_{call.source}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        stored_filename = f"{uuid.uuid4().hex}.docx"

        tenant_folder = os.path.join(
            current_app.config.get('UPLOAD_FOLDER', 'uploads'),
            str(project.tenant_id or 'global')
        )
        os.makedirs(tenant_folder, exist_ok=True)

        file_path = os.path.join(tenant_folder, stored_filename)
        with open(file_path, 'wb') as f:
            f.write(docx_content)

        # Create document record
        doc_record = ProjectDocument(
            tenant_id=project.tenant_id,
            project_id=project.id,
            filename=filename,
            stored_filename=stored_filename,
            file_type='docx',
            file_size=len(docx_content),
            description=f"Proposta para {call.source} gerada por IA em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            uploaded_by_id=user_id
        )

        db.session.add(doc_record)
        db.session.commit()

        logger.info(f"Saved proposal Word document: {filename} for project {project.code}")
        return doc_record

    except Exception as e:
        logger.error(f"Error saving proposal Word document: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return None


def _add_risks_content(story, content, styles, section_style):
    """Add risk analysis content to PDF."""
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    if not isinstance(content, dict):
        story.append(Paragraph(str(content), styles['Normal']))
        return

    # Overall risk level and score
    nivel_risco = content.get('nivel_risco_geral', 'N/A')
    score_risco = content.get('score_risco', 'N/A')

    risk_color = '#dc3545' if nivel_risco in ['alto', 'crítico'] else '#ffc107' if nivel_risco == 'médio' else '#28a745'

    story.append(Paragraph("Resumo de Risco", section_style))
    story.append(Paragraph(f"<b>Nível de Risco Geral:</b> <font color='{risk_color}'>{nivel_risco.upper()}</font>", styles['Normal']))
    story.append(Paragraph(f"<b>Score de Risco:</b> {score_risco}/100", styles['Normal']))
    story.append(Spacer(1, 15))

    # Burn rate analysis
    burn_rate = content.get('burn_rate', {})
    if burn_rate:
        story.append(Paragraph("Análise de Burn Rate", section_style))
        story.append(Paragraph(f"<b>Gasto Mensal Estimado:</b> R$ {burn_rate.get('mensal', 0):,.2f}", styles['Normal']))
        story.append(Paragraph(f"<b>Projeção de Gasto Final:</b> R$ {burn_rate.get('projecao_final', 0):,.2f}", styles['Normal']))
        status_burn = burn_rate.get('status', 'N/A')
        burn_color = '#dc3545' if 'acima' in status_burn else '#ffc107' if 'risco' in status_burn else '#28a745'
        story.append(Paragraph(f"<b>Status:</b> <font color='{burn_color}'>{status_burn}</font>", styles['Normal']))
        story.append(Spacer(1, 15))

    # Completion forecast
    previsao = content.get('previsao_conclusao', {})
    if previsao:
        story.append(Paragraph("Previsão de Conclusão", section_style))
        status_prev = previsao.get('status', 'N/A')
        prev_color = '#dc3545' if 'atrasado' in status_prev else '#ffc107' if 'risco' in status_prev else '#28a745'
        story.append(Paragraph(f"<b>Status:</b> <font color='{prev_color}'>{status_prev}</font>", styles['Normal']))
        story.append(Paragraph(f"<b>Estimativa:</b> {previsao.get('estimativa', 'N/A')}", styles['Normal']))
        story.append(Paragraph(f"<b>Confiança:</b> {previsao.get('confianca', 'N/A')}", styles['Normal']))
        story.append(Spacer(1, 15))

    # Alerts
    alertas = content.get('alertas', [])
    if alertas:
        story.append(Paragraph("⚠️ Alertas Imediatos", section_style))
        for alerta in alertas:
            story.append(Paragraph(f"• <font color='#dc3545'>{alerta}</font>", styles['Normal']))
        story.append(Spacer(1, 15))

    # Identified risks
    riscos = content.get('riscos_identificados', [])
    if riscos:
        story.append(Paragraph(f"Riscos Identificados ({len(riscos)})", section_style))

        for i, risk in enumerate(riscos, 1):
            categoria = risk.get('categoria', 'N/A')
            severidade = risk.get('severidade', 'N/A')
            sev_color = '#dc3545' if severidade in ['alta', 'crítica'] else '#ffc107' if severidade == 'média' else '#17a2b8'

            story.append(Paragraph(f"<b>{i}. [{categoria.upper()}] - Severidade: <font color='{sev_color}'>{severidade.upper()}</font></b>", styles['Normal']))
            story.append(Paragraph(f"<b>Descrição:</b> {risk.get('descricao', 'N/A')}", styles['Normal']))
            story.append(Paragraph(f"<b>Probabilidade:</b> {risk.get('probabilidade', 'N/A')}", styles['Normal']))
            story.append(Paragraph(f"<b>Impacto:</b> {risk.get('impacto', 'N/A')}", styles['Normal']))
            story.append(Paragraph(f"<b>Mitigação:</b> {risk.get('mitigacao', 'N/A')}", styles['Normal']))
            story.append(Spacer(1, 10))

    # Recommendations
    recomendacoes = content.get('recomendacoes', [])
    if recomendacoes:
        story.append(Paragraph("💡 Recomendações", section_style))
        for rec in recomendacoes:
            story.append(Paragraph(f"• {rec}", styles['Normal']))
        story.append(Spacer(1, 10))


def _add_report_content(story, content, styles, section_style):
    """Add report content to PDF."""
    from reportlab.platypus import Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle

    # Get the actual report text
    report_text = content
    if isinstance(content, dict):
        report_text = content.get('report', str(content))

    if isinstance(report_text, str):
        # Parse markdown-like content
        lines = report_text.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                story.append(Spacer(1, 8))
            elif line.startswith('### '):
                story.append(Spacer(1, 10))
                story.append(Paragraph(f"<b>{line[4:]}</b>", styles['Normal']))
            elif line.startswith('## '):
                story.append(Spacer(1, 12))
                story.append(Paragraph(line[3:], section_style))
            elif line.startswith('# '):
                story.append(Spacer(1, 15))
                story.append(Paragraph(line[2:], section_style))
            elif line.startswith('- ') or line.startswith('* '):
                story.append(Paragraph(f"• {line[2:]}", styles['Normal']))
            elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
                story.append(Paragraph(line, styles['Normal']))
            else:
                # Clean markdown formatting
                clean_line = line.replace('**', '').replace('__', '').replace('*', '').replace('_', '')
                if clean_line:
                    story.append(Paragraph(clean_line, styles['Normal']))
    else:
        story.append(Paragraph(str(report_text), styles['Normal']))


def _add_matching_content(story, content, styles, section_style):
    """Add matching results content to PDF."""
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    matches = content.get('matches', []) if isinstance(content, dict) else content

    if matches:
        story.append(Paragraph(f"Editais Compatíveis ({len(matches)})", section_style))
        story.append(Spacer(1, 10))

        for i, match in enumerate(matches, 1):
            score = match.get('score', match.get('compatibility_score', 0))
            score_color = '#28a745' if score >= 80 else '#ffc107' if score >= 60 else '#dc3545'

            story.append(Paragraph(f"<b>{i}. {match.get('call_title', match.get('title', 'Edital'))}</b>", styles['Normal']))
            story.append(Paragraph(f"<b>Fonte:</b> {match.get('call_source', match.get('source', 'N/A'))}", styles['Normal']))
            story.append(Paragraph(f"<b>Compatibilidade:</b> <font color='{score_color}'><b>{score}%</b></font>", styles['Normal']))

            # Reasons/justification
            reasons = match.get('reasons', match.get('justificativa', match.get('pontos_fortes', [])))
            if reasons:
                if isinstance(reasons, str):
                    reasons = [reasons]
                story.append(Paragraph("<b>Pontos de Compatibilidade:</b>", styles['Normal']))
                for reason in reasons:
                    story.append(Paragraph(f"  ✓ {reason}", styles['Normal']))

            # Gaps/weaknesses
            gaps = match.get('gaps', match.get('pontos_fracos', []))
            if gaps:
                if isinstance(gaps, str):
                    gaps = [gaps]
                story.append(Paragraph("<b>Pontos de Atenção:</b>", styles['Normal']))
                for gap in gaps:
                    story.append(Paragraph(f"  ⚠ {gap}", styles['Normal']))

            # Recommendations
            recomendacoes = match.get('recomendacoes', match.get('recommendations', []))
            if recomendacoes:
                if isinstance(recomendacoes, str):
                    recomendacoes = [recomendacoes]
                story.append(Paragraph("<b>Recomendações:</b>", styles['Normal']))
                for rec in recomendacoes:
                    story.append(Paragraph(f"  → {rec}", styles['Normal']))

            deadline = match.get('call_deadline', match.get('deadline', ''))
            if deadline:
                story.append(Paragraph(f"<b>Prazo:</b> {deadline}", styles['Normal']))

            url = match.get('call_url', match.get('url', ''))
            if url:
                story.append(Paragraph(f"<b>Link:</b> {url}", styles['Normal']))

            story.append(Spacer(1, 15))
    else:
        story.append(Paragraph("Nenhum edital compatível encontrado.", styles['Normal']))


def _add_document_analysis_content(story, content, styles, section_style):
    """Add document analysis content to PDF."""
    from reportlab.platypus import Paragraph, Spacer

    if not isinstance(content, dict):
        story.append(Paragraph(str(content), styles['Normal']))
        return

    # Summary / resumo
    resumo = content.get('summary', content.get('resumo', content.get('resumo_executivo', '')))
    if resumo:
        story.append(Paragraph("Resumo", section_style))
        story.append(Paragraph(resumo, styles['Normal']))
        story.append(Spacer(1, 15))

    # Objective / objetivo
    objetivo = content.get('objetivo', content.get('objective', ''))
    if objetivo:
        story.append(Paragraph("Objetivo", section_style))
        story.append(Paragraph(objetivo, styles['Normal']))
        story.append(Spacer(1, 15))

    # Key points / pontos principais
    pontos = content.get('key_points', content.get('pontos_principais', content.get('pontos_chave', [])))
    if pontos:
        story.append(Paragraph("Pontos Principais", section_style))
        if isinstance(pontos, str):
            pontos = [pontos]
        for point in pontos:
            story.append(Paragraph(f"• {point}", styles['Normal']))
        story.append(Spacer(1, 15))

    # Requirements / requisitos
    requisitos = content.get('requirements', content.get('requisitos', content.get('requisitos_tecnicos', [])))
    if requisitos:
        story.append(Paragraph("Requisitos Identificados", section_style))
        if isinstance(requisitos, str):
            requisitos = [requisitos]
        for req in requisitos:
            story.append(Paragraph(f"• {req}", styles['Normal']))
        story.append(Spacer(1, 15))

    # Elegibility / elegibilidade
    elegibilidade = content.get('elegibilidade', content.get('eligibility', content.get('publico_alvo', '')))
    if elegibilidade:
        story.append(Paragraph("Elegibilidade / Público-Alvo", section_style))
        if isinstance(elegibilidade, list):
            for item in elegibilidade:
                story.append(Paragraph(f"• {item}", styles['Normal']))
        else:
            story.append(Paragraph(elegibilidade, styles['Normal']))
        story.append(Spacer(1, 15))

    # Funding / recursos
    recursos = content.get('recursos', content.get('funding', content.get('valor_maximo', '')))
    if recursos:
        story.append(Paragraph("Recursos / Financiamento", section_style))
        story.append(Paragraph(str(recursos), styles['Normal']))
        story.append(Spacer(1, 15))

    # Deadlines / prazos
    prazos = content.get('prazos', content.get('deadlines', content.get('cronograma', [])))
    if prazos:
        story.append(Paragraph("Prazos", section_style))
        if isinstance(prazos, list):
            for prazo in prazos:
                story.append(Paragraph(f"• {prazo}", styles['Normal']))
        else:
            story.append(Paragraph(str(prazos), styles['Normal']))
        story.append(Spacer(1, 15))

    # Analysis / análise detalhada
    analise = content.get('analysis', content.get('analise', content.get('analise_detalhada', '')))
    if analise:
        story.append(Paragraph("Análise Detalhada", section_style))
        story.append(Paragraph(analise, styles['Normal']))
        story.append(Spacer(1, 15))

    # Recommendations
    recomendacoes = content.get('recomendacoes', content.get('recommendations', []))
    if recomendacoes:
        story.append(Paragraph("Recomendações", section_style))
        if isinstance(recomendacoes, str):
            recomendacoes = [recomendacoes]
        for rec in recomendacoes:
            story.append(Paragraph(f"• {rec}", styles['Normal']))
        story.append(Spacer(1, 15))

    # If content has other keys not handled, add them
    handled_keys = {'summary', 'resumo', 'resumo_executivo', 'objetivo', 'objective',
                    'key_points', 'pontos_principais', 'pontos_chave', 'requirements',
                    'requisitos', 'requisitos_tecnicos', 'elegibilidade', 'eligibility',
                    'publico_alvo', 'recursos', 'funding', 'valor_maximo', 'prazos',
                    'deadlines', 'cronograma', 'analysis', 'analise', 'analise_detalhada',
                    'recomendacoes', 'recommendations', 'document_saved', 'document_id', 'document_name'}

    for key, value in content.items():
        if key not in handled_keys and value:
            story.append(Paragraph(f"<b>{key.replace('_', ' ').title()}</b>", styles['Normal']))
            if isinstance(value, list):
                for item in value:
                    story.append(Paragraph(f"• {item}", styles['Normal']))
            else:
                story.append(Paragraph(str(value), styles['Normal']))
            story.append(Spacer(1, 10))

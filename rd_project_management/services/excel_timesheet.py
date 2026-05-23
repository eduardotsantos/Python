"""Service for Excel timesheet import/export."""

from io import BytesIO
from datetime import datetime, date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def create_timesheet_template(project, milestones, resources):
    """Create an Excel template for timesheet import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    headers = ["Data (DD/MM/AAAA)", "Horas", "Atividade", "Marco (opcional)", "Recurso (opcional)", "Observações (opcional)"]
    col_widths = [18, 10, 40, 30, 30, 40]

    for col, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = width

    for row in range(2, 52):
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            if col == 1:
                cell.number_format = 'DD/MM/YYYY'
            elif col == 2:
                cell.number_format = '0.0'

    ws_info = wb.create_sheet("Informações")
    ws_info.sheet_properties.tabColor = "92D050"

    ws_info['A1'] = "INSTRUÇÕES DE PREENCHIMENTO"
    ws_info['A1'].font = Font(bold=True, size=14)
    ws_info.merge_cells('A1:D1')

    instructions = [
        "",
        "1. Data: Formato DD/MM/AAAA (ex: 23/05/2026)",
        "2. Horas: Número decimal (ex: 8 ou 2.5)",
        "3. Atividade: Descrição obrigatória da atividade realizada",
        "4. Marco: Selecione da lista ou deixe em branco",
        "5. Recurso: Selecione da lista ou deixe em branco",
        "6. Observações: Campo livre opcional",
        "",
        f"Projeto: {project.title}",
        f"Data de geração: {date.today().strftime('%d/%m/%Y')}",
    ]

    for i, text in enumerate(instructions, 2):
        ws_info[f'A{i}'] = text

    ws_info.column_dimensions['A'].width = 60

    if milestones or resources:
        ws_listas = wb.create_sheet("Listas")
        ws_listas.sheet_properties.tabColor = "FFC000"

        ws_listas['A1'] = "Marcos"
        ws_listas['A1'].font = Font(bold=True)
        milestone_names = [m.title for m in milestones] if milestones else []
        for i, name in enumerate(milestone_names, 2):
            ws_listas[f'A{i}'] = name

        ws_listas['B1'] = "Recursos"
        ws_listas['B1'].font = Font(bold=True)
        resource_names = [r.name for r in resources] if resources else []
        for i, name in enumerate(resource_names, 2):
            ws_listas[f'B{i}'] = name

        ws_listas.column_dimensions['A'].width = 40
        ws_listas.column_dimensions['B'].width = 40

        if milestone_names:
            milestone_range = f"Listas!$A$2:$A${len(milestone_names)+1}"
            dv_milestone = DataValidation(type="list", formula1=milestone_range, allow_blank=True)
            dv_milestone.error = "Selecione um marco da lista"
            dv_milestone.errorTitle = "Marco inválido"
            ws.add_data_validation(dv_milestone)
            dv_milestone.add(f'D2:D51')

        if resource_names:
            resource_range = f"Listas!$B$2:$B${len(resource_names)+1}"
            dv_resource = DataValidation(type="list", formula1=resource_range, allow_blank=True)
            dv_resource.error = "Selecione um recurso da lista"
            dv_resource.errorTitle = "Recurso inválido"
            ws.add_data_validation(dv_resource)
            dv_resource.add(f'E2:E51')

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_timesheet_excel(file_stream, project_id, tenant_id, user_id, milestones, resources):
    """Parse an Excel file and return timesheet entries."""
    wb = load_workbook(file_stream)
    ws = wb.active

    milestone_map = {m.title.lower(): m.id for m in milestones} if milestones else {}
    resource_map = {r.name.lower(): r.id for r in resources} if resources else {}

    entries = []
    errors = []

    for row_num in range(2, ws.max_row + 1):
        date_val = ws.cell(row=row_num, column=1).value
        hours_val = ws.cell(row=row_num, column=2).value
        activity_val = ws.cell(row=row_num, column=3).value
        milestone_val = ws.cell(row=row_num, column=4).value
        resource_val = ws.cell(row=row_num, column=5).value
        notes_val = ws.cell(row=row_num, column=6).value

        if not date_val and not hours_val and not activity_val:
            continue

        if not date_val:
            errors.append(f"Linha {row_num}: Data é obrigatória")
            continue

        if not hours_val:
            errors.append(f"Linha {row_num}: Horas é obrigatório")
            continue

        if not activity_val:
            errors.append(f"Linha {row_num}: Atividade é obrigatória")
            continue

        try:
            if isinstance(date_val, datetime):
                entry_date = date_val.date()
            elif isinstance(date_val, date):
                entry_date = date_val
            elif isinstance(date_val, str):
                date_val = date_val.strip()
                for fmt in ['%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y']:
                    try:
                        entry_date = datetime.strptime(date_val, fmt).date()
                        break
                    except ValueError:
                        continue
                else:
                    errors.append(f"Linha {row_num}: Data inválida '{date_val}'")
                    continue
            else:
                errors.append(f"Linha {row_num}: Formato de data não reconhecido")
                continue
        except Exception as e:
            errors.append(f"Linha {row_num}: Erro ao processar data - {str(e)}")
            continue

        try:
            hours = float(hours_val)
            if hours <= 0 or hours > 24:
                errors.append(f"Linha {row_num}: Horas deve ser entre 0 e 24")
                continue
        except (ValueError, TypeError):
            errors.append(f"Linha {row_num}: Horas inválidas '{hours_val}'")
            continue

        activity = str(activity_val).strip()
        if len(activity) < 3:
            errors.append(f"Linha {row_num}: Atividade muito curta")
            continue

        milestone_id = None
        if milestone_val:
            milestone_key = str(milestone_val).strip().lower()
            milestone_id = milestone_map.get(milestone_key)

        resource_id = None
        if resource_val:
            resource_key = str(resource_val).strip().lower()
            resource_id = resource_map.get(resource_key)

        notes = str(notes_val).strip() if notes_val else None

        entries.append({
            'tenant_id': tenant_id,
            'project_id': project_id,
            'user_id': user_id,
            'date': entry_date,
            'hours': hours,
            'activity': activity,
            'milestone_id': milestone_id,
            'resource_id': resource_id,
            'notes': notes
        })

    return entries, errors
